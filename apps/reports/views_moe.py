from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from apps.accounts.models import UserRole
from apps.platform_management.decorators import school_context_required


def _get_moe_data(school, current_ay):
    """Aggregates all statistics needed for MoE report."""
    from apps.students.models import StudentProfile
    from apps.enrollment.models import StudentEnrollment, EnrollmentStatus
    from apps.teachers.models import TeacherProfile
    from apps.academics.models import Grade, Section

    data = {}

    students_qs = StudentProfile.objects.filter(school=school)
    data['total_students'] = students_qs.count()
    data['male_students'] = students_qs.filter(gender='M').count()
    data['female_students'] = students_qs.filter(gender='F').count()

    teachers_qs = TeacherProfile.objects.filter(school=school)
    data['total_teachers'] = teachers_qs.count()
    data['male_teachers'] = teachers_qs.filter(gender='M').count()
    data['female_teachers'] = teachers_qs.filter(gender='F').count()

    if current_ay:
        enrollments = StudentEnrollment.objects.filter(
            school=school, academic_year=current_ay, status=EnrollmentStatus.ACTIVE
        )
        data['enrolled_students'] = enrollments.count()

        # Grade-wise breakdown
        grade_breakdown = []
        for grade in Grade.objects.filter(school=school).order_by('name'):
            grade_enr = enrollments.filter(grade=grade)
            grade_breakdown.append({
                'grade': grade.name,
                'total': grade_enr.count(),
                'male': grade_enr.filter(student__gender='M').count(),
                'female': grade_enr.filter(student__gender='F').count(),
            })
        data['grade_breakdown'] = grade_breakdown
    else:
        data['enrolled_students'] = 0
        data['grade_breakdown'] = []

    data['current_ay'] = current_ay
    data['school'] = school
    return data


@login_required
@school_context_required
def moe_report_view(request):
    """MoE statistical dashboard view."""
    school = getattr(request, 'school', None)
    user = request.user

    if user.role not in [UserRole.SCHOOL_ADMIN, UserRole.PRINCIPAL, UserRole.REGISTRAR, UserRole.SUPER_ADMIN]:
        from django.contrib import messages
        from django.shortcuts import redirect
        messages.error(request, "Unauthorized.")
        return redirect('index')

    current_ay = getattr(request, 'academic_year', None)
    moe_data = _get_moe_data(school, current_ay)
    return render(request, 'reports/moe_report.html', {'moe_data': moe_data})


@login_required
def moe_export_excel_view(request):
    """Exports MoE statistics as a multi-sheet Excel workbook."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from django.utils import timezone

    school = getattr(request, 'school', None)
    current_ay = getattr(request, 'academic_year', None)
    moe_data = _get_moe_data(school, current_ay)

    wb = openpyxl.Workbook()

    # ── Sheet 1: Summary ──────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Summary"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")

    ws1.append(["Ministry of Education — School Statistics Report"])
    ws1['A1'].font = Font(bold=True, size=14)
    ws1.append([f"School: {school.name}"])
    ay_name = current_ay.name if current_ay else "N/A"
    ws1.append([f"Academic Year: {ay_name}"])
    ws1.append([f"Generated On: {timezone.now().strftime('%Y-%m-%d %H:%M')}"])
    ws1.append([])

    headers = ["Metric", "Value"]
    ws1.append(headers)
    for cell in ws1[ws1.max_row]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    rows = [
        ["Total Students", moe_data['total_students']],
        ["Male Students", moe_data['male_students']],
        ["Female Students", moe_data['female_students']],
        ["Currently Enrolled", moe_data['enrolled_students']],
        ["Total Teachers", moe_data['total_teachers']],
    ]
    for row in rows:
        ws1.append(row)

    ws1.column_dimensions['A'].width = 30
    ws1.column_dimensions['B'].width = 20

    # ── Sheet 2: Grade-wise Enrollment ──────────────────────────────
    ws2 = wb.create_sheet("Grade Enrollment")
    ws2.append(["Grade", "Total", "Male", "Female"])
    for cell in ws2[1]:
        cell.font = header_font
        cell.fill = header_fill
    for g in moe_data.get('grade_breakdown', []):
        ws2.append([g['grade'], g['total'], g['male'], g['female']])
    ws2.column_dimensions['A'].width = 20
    ws2.column_dimensions['B'].width = 15
    ws2.column_dimensions['C'].width = 15
    ws2.column_dimensions['D'].width = 15

    # Write response
    from io import BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"MoE_Report_{school.code}_{ay_name.replace(' ', '_')}.xlsx"
    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
